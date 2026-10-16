"""Bounded, local-only extraction into reviewable school-document drafts.

No database writes, external AI calls, macro execution or automatic publication.
"""

import csv
from io import BytesIO, StringIO
from pathlib import Path
import re
import time
import zipfile
import xml.etree.ElementTree as ET


MAX_BYTES = 8 * 1024 * 1024
MAX_ROWS = 1000
MAX_TEXT = 150000
DAYS = {name[:3].lower(): name for name in ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]}
SUBJECTS = {
    "acc": "Accountancy", "accounts": "Accountancy", "accountancy": "Accountancy",
    "bs": "Business Studies", "business studies": "Business Studies", "eco": "Economics", "economics": "Economics",
    "eng": "English", "english": "English", "com": "Computer Science", "computer": "Computer Science",
    "computer science": "Computer Science", "pe": "Physical Education", "physical education": "Physical Education",
    "math": "Mathematics", "math pcmb": "Mathematics", "math pcmb pcm": "Mathematics", "mathematics": "Mathematics",
    "fine arts": "Fine Arts", "psychology": "Psychology", "compulsory arabic": "Compulsory Arabic",
    "c arabic": "Compulsory Arabic", "islamic culture": "Islamic Culture",
}


class ImportProblem(ValueError):
    pass


def clean(value):
    return str("" if value is None else value).replace("\x00", "").strip()


def subject_label(value):
    key = re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()
    compact = re.sub(r"[^a-z0-9]", "", value.lower())
    return SUBJECTS.get(key, SUBJECTS.get(compact, value))


def checked_zip(content):
    archive = zipfile.ZipFile(BytesIO(content))
    entries = archive.infolist()
    if len(entries) > 2000 or sum(entry.file_size for entry in entries) > 30 * 1024 * 1024:
        archive.close()
        raise ImportProblem("The expanded document is too large. Split it into smaller files.")
    return archive


def extract_source(filename, content):
    """Return local text/tables. Inputs are data, never instructions."""
    if not content or len(content) > MAX_BYTES:
        raise ImportProblem("Choose a non-empty file of at most 8 MB.")
    suffix = Path(filename).suffix.lower()
    tables, text, warnings = [], "", []
    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ImportProblem("Save the CSV using UTF-8 encoding and try again.") from None
        try:
            dialect = csv.Sniffer().sniff(text[:8000], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        grid = []
        for row in csv.reader(StringIO(text), dialect):
            if len(grid) > MAX_ROWS or len(row) > 100:
                raise ImportProblem("This sheet exceeds 1,000 rows or 100 columns. Split it first.")
            if any(cell.strip() for cell in row):
                grid.append(row)
        tables = [grid]
    elif suffix == ".xlsx":
        with checked_zip(content):
            pass
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            if len(book.worksheets) > 10:
                raise ImportProblem("Choose a workbook with at most 10 sheets.")
            count = 0
            for sheet in book.worksheets:
                grid = []
                for row in sheet.iter_rows(values_only=True):
                    count += 1
                    if count > MAX_ROWS or len(row) > 100:
                        raise ImportProblem("This workbook exceeds 1,000 rows or 100 columns. Split it first.")
                    if any(value is not None for value in row):
                        grid.append([clean(value) for value in row])
                if grid:
                    tables.append(grid)
                    text += f"\nSheet: {sheet.title}\n" + "\n".join(" | ".join(row) for row in grid)
            warnings.append("Formula cells use saved values only. Check any blank or out-of-date results.")
        finally:
            book.close()
    elif suffix == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise ImportProblem("This file does not contain a valid PDF header.")
        import pdfplumber
        ocr_deadline = time.monotonic() + 120
        with pdfplumber.open(BytesIO(content)) as pdf:
            if len(pdf.pages) > 20:
                raise ImportProblem("Upload at most 20 PDF pages at a time.")
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                if not page_text.strip():
                    if len(pdf.pages) > 5:
                        raise ImportProblem("Scanned PDF OCR is limited to 5 pages. Split this document.")
                    from school_ocr import extract_image
                    buffer = BytesIO()
                    page.to_image(resolution=130).original.save(buffer, format="PNG")
                    scanned = extract_image(buffer.getvalue(), deadline=ocr_deadline)
                    page_text = scanned["text"]
                    tables.extend(scanned["tables"])
                    warnings.extend(scanned["warnings"])
                text += "\n" + page_text
                for grid in page.extract_tables():
                    tables.append([[clean(cell) for cell in row] for row in grid])
                if len(text) > MAX_TEXT or sum(len(grid) for grid in tables) > MAX_ROWS:
                    raise ImportProblem("The extracted document is too large. Split it first.")
        if not text.strip():
            raise ImportProblem("No readable text was found. Upload a clearer original or enter verified rows manually.")
    elif suffix == ".docx":
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        with checked_zip(content) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
        text = "\n".join("".join(node.itertext()) for node in root.findall(".//w:t", ns))
        for table in root.findall(".//w:tbl", ns):
            tables.append([["\n".join(node.text or "" for node in cell.findall(".//w:t", ns))
                            for cell in row.findall("w:tc", ns)] for row in table.findall("w:tr", ns)])
    elif suffix in {".jpg", ".jpeg", ".png", ".webp"}:
        from school_ocr import extract_image
        return extract_image(content)
    else:
        raise ImportProblem("Supported draft sources: CSV, XLSX, PDF, DOCX, JPG, PNG and WebP.")
    if len(text) > MAX_TEXT or sum(len(grid) for grid in tables) > MAX_ROWS:
        raise ImportProblem("The extracted document exceeds the safe review size. Split it first.")
    return dict(text=text.strip(), tables=tables, warnings=warnings)


TIMETABLE_HEADERS = {
    "day": {"day", "weekday"}, "period": {"period", "period no", "period number"},
    "start": {"start", "start time", "from"}, "end": {"end", "end time", "to"},
    "activity": {"subject", "activity", "lesson"}, "teacher": {"teacher", "staff", "teacher name"},
    "class": {"class", "division", "class division"},
}


def timetable_rows(tables):
    output, warnings = [], []
    for table in tables:
        for index, header in enumerate(table):
            mapping = {}
            for col, value in enumerate(header):
                label = re.sub(r"[_\W]+", " ", clean(value).lower()).strip()
                for key, variants in TIMETABLE_HEADERS.items():
                    if label in variants:
                        mapping[key] = col
            if {"day", "period", "activity"} <= mapping.keys():
                for source in table[index + 1:]:
                    row = {key: clean(source[col]) if col < len(source) else "" for key, col in mapping.items()}
                    day = DAYS.get(row["day"][:3].lower())
                    if day:
                        output.append({**dict.fromkeys(TIMETABLE_HEADERS, ""), **row, "day": day,
                                       "activity": subject_label(row["activity"])})
                break
            # Weekly grids: a header immediately followed by a weekday row.
            if index + 1 < len(table) and clean(table[index + 1][0])[:3].lower() in DAYS:
                for source in table[index + 1:]:
                    if not source:
                        continue
                    day = DAYS.get(clean(source[0])[:3].lower())
                    if not day:
                        continue
                    for col in range(1, min(len(header), len(source))):
                        heading, activity = clean(header[col]), clean(source[col])
                        if not heading or (not activity and 'break' not in heading.lower()):
                            continue
                        period = re.search(r"(?:^|\bperiod\s*)(\d{1,2})(?![\d:.])\b", heading, re.I)
                        times = re.findall(r"\b\d{1,2}[:.]\d{2}\b", heading)
                        if "break" in heading.lower():
                            activity = "Break"
                            break_number = re.search(r'break[ _-]*(\d+)', heading, re.I)
                            period_value = 'Break ' + str(int(break_number[1])) if break_number else ''
                        else:
                            period_value = period[1] if period else ''
                        output.append(dict(day=day, period=period_value, start=times[0] if times else "",
                                           end=times[1] if len(times) > 1 else "", activity=subject_label(activity), teacher="", **{"class": ""}))
                warnings.append("Grid cells are preserved as extracted. Review combined subjects, teacher names, periods and AM/PM times before approval.")
                break
    return output, warnings


def syllabus_rows(tables):
    output, current_subject, section = [], "", ""
    for table in tables:
        for row in table:
            if len(row) != 2:
                continue
            left, right = clean(row[0]), clean(row[1])
            if left.lower() == "subject" or not right or left.lower().startswith("approved by"):
                continue
            if left:
                # A merged header is not a subject. Keep unknown short subject labels for review.
                if len(left) > 70 or "academic year" in left.lower() or "syllabus" in left.lower():
                    continue
                next_subject = subject_label(left)
                if next_subject != current_subject:
                    section = ""
                current_subject = next_subject
            if not current_subject:
                continue
            for line in right.splitlines():
                line = clean(line)
                if not line:
                    continue
                match = re.match(r"(?:chapter|ch\.?)[\s:]*(\d+)\s*[:.\-]?\s*(.*)", line, re.I)
                if match:
                    output.append(dict(subject=current_subject, section=section, chapter=match[1], topic=match[2] or line))
                elif line.endswith(":") or re.match(r"^(?:part\s+[IVX]+\b|writing skill|literature|flamingo|vistas|introductory macroeconomics|indian economic development)", line, re.I):
                    section = line.rstrip(":")
                elif output and output[-1]["subject"] == current_subject and output[-1]["chapter"] and output[-1]["section"] == section:
                    output[-1]["topic"] += " " + line
                else:
                    output.append(dict(subject=current_subject, section=section, chapter="", topic=re.sub(r"^\d+[.)]\s*", "", line)))
    return output


def make_draft(kind, filename, content):
    if kind not in {"timetable", "teacher_timetable", "syllabus"}:
        raise ImportProblem("Choose student timetable, teacher timetable or syllabus.")
    source = extract_source(filename, content)
    if kind in {"timetable", "teacher_timetable"}:
        rows, warnings = timetable_rows(source["tables"])
    else:
        rows, warnings = syllabus_rows(source["tables"]), []
    if len(rows) > MAX_ROWS:
        raise ImportProblem("This document has too many extracted entries. Split it first.")
    if not rows:
        warnings.append("No reliable table structure was detected. The source text is available for review; no records have been inferred.")
    warnings.append("Draft only. Review every extracted entry and confirm class, year and assessment before publication.")
    year = re.search(r"(?:academic\s*year\s*[:\-]?\s*)(20\d{2}\s*[-–]\s*\d{2,4})", source["text"], re.I)
    term = re.search(r"\bterm\s*:\s*([IVX]+|\d+)\b", source["text"], re.I)
    grade = re.search(r"\bgrade\s*:\s*([IVX]+|\d+)\s*(?:\(([^)]+)\))?", source["text"], re.I)
    return dict(kind=kind, filename=Path(filename.replace("\\", "/")).name[:180],
                rows=rows, source_text=source["text"], warnings=source["warnings"] + warnings,
                suggested_scope=dict(year=year[1].replace(" ", "") if year else "", grade=grade[1] if grade else "",
                                     stream=grade[2] or "" if grade else "", assessment=f"Term {term[1]}" if term else ""))
